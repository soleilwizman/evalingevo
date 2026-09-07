from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from pathlib import Path
D=Document()
s=D.sections[0]; s.top_margin=Inches(.7); s.bottom_margin=Inches(.65); s.left_margin=s.right_margin=Inches(.8)
s.page_width=Inches(8.5); s.page_height=Inches(11)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
 st=D.styles[name]; st.font.name='Calibri'; st.font.color.rgb=RGBColor(0,0,0)
D.styles['Normal'].font.size=Pt(11)
D.styles['Normal'].paragraph_format.space_after=Pt(7)
D.styles['Normal'].paragraph_format.line_spacing=1.08
for name,size in [('Title',25),('Heading 1',19),('Heading 2',13),('Heading 3',11)]:
 D.styles[name].font.size=Pt(size)
 D.styles[name].paragraph_format.space_before=Pt(10)
 D.styles[name].paragraph_format.space_after=Pt(7)
D.styles['Subtitle'].font.size=Pt(11)
foot=s.footer.paragraphs[0]; foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=foot.add_run(); fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); r._r.addnext(fld)
def p(t): D.add_paragraph(t)
def h(t): D.add_heading(t,2)
def page(t): D.add_page_break(); D.add_heading(t,1)
def table(headers, rows, widths):
 t=D.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
 for col,w in zip(t.columns,widths): col.width=Inches(w)
 for c,w,txt in zip(t.rows[0].cells,widths,headers): c.width=Inches(w); c.text=txt
 for row in rows:
  for c,w,txt in zip(t.add_row().cells,widths,row): c.width=Inches(w); c.text=txt
 for i,row in enumerate(t.rows):
  for c in row.cells:
   c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=c._tc.get_or_add_tcPr(); borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    b=OxmlElement('w:'+edge); b.set(qn('w:val'),'single'); b.set(qn('w:sz'),'4'); b.set(qn('w:color'),'D9D9D9'); borders.append(b)
   pr.append(borders); margins=OxmlElement('w:tcMar')
   for edge in ['top','left','bottom','right']:
    b=OxmlElement('w:'+edge); b.set(qn('w:w'),'90'); b.set(qn('w:type'),'dxa'); margins.append(b)
   pr.append(margins)
   if i==0:
    sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'E8EEF3'); pr.append(sh)
   for pp in c.paragraphs:
    pp.paragraph_format.space_after=Pt(2); pp.paragraph_format.line_spacing=1.0
    for run in pp.runs: run.font.size=Pt(10); run.bold=(i==0)
  if i==0:
   rep=OxmlElement('w:tblHeader'); row._tr.get_or_add_trPr().append(rep)
  no=OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(no)
 D.add_paragraph().paragraph_format.space_after=Pt(0)

D.add_heading('Evo 2 and regulatory interactions',0)
D.add_paragraph('Project proposal and MVP report\nSeptember 5 2026, numbers revised September 7 2026 after the recoded pipeline rerun',style='Subtitle')
D.add_heading('Part 1  The project proposal',1)
p('I propose to test whether Evo 2 can help predict and explain how pairs of human regulatory variants interact. The project connects three levels of evidence: prediction of measured effects, identification of useful internal model features, and experimental tests of the biological hypotheses those features suggest.')
p('Today we built the first step: a reproducible comparison between a frozen Evo 2 sequence score and published measurements of two-variant interactions. The result is a narrow null: this particular score shows essentially no association with measured interactions in the selected K562 assay. That result defines the next experiments. It does not establish whether other model readouts, internal representations, or cell-specific predictors will succeed.')
h('The biological question')
p('A regulatory DNA sequence helps control gene expression. Two nearby variants may each change its activity, and their combined effect may differ from what their separate effects predict. We call that departure epistasis. For example, two changes might weaken the same binding site, so the second change has little additional effect after the first. Alternatively, they may disrupt two cooperating sites. These are hypotheses about regulatory logic; an interaction measurement alone does not tell us which mechanism produced it.')
h('The computational question')
p('Evo 2 learns to predict DNA bases from sequence context. Its output assigns a likelihood to a sequence. A sequence that looks more plausible to the model is not automatically a sequence that drives stronger expression in a particular human cell. The proposal therefore asks whether regulatory information is accessible from the output score, recoverable from internal representations, or absent under the tested conditions. Evo 2 supports sequence scoring and access to intermediate embeddings, making these separate tests feasible. [2]')
h('What success would mean')
p('The eventual contribution would be a benchmark with credible measurement uncertainty, a predictor that generalizes to unseen regulatory regions, and experimentally supported explanations for a subset of interactions. A strong negative study is also possible if careful measurement and independent replication establish where the approach fails. We should choose the paper’s claims after these tests, rather than assume in advance that Evo contains a discoverable regulatory mechanism.')

page('Part 2  What we built today')
h('Start with four versions of the same DNA fragment')
p('Our MVP, or minimum viable experiment, uses released data from Siraj and colleagues. We performed computational reconstruction and analysis of their measurements; we did not run a new wet-lab assay today. The source study and data release are listed in references 1a and 1b.')
p('A massively parallel reporter assay, or MPRA, tests many regulatory sequences using reporter constructs. RNA output is compared with DNA abundance to estimate regulatory activity. For each variant pair, we need four matched sequence versions, often called a quartet. Keeping the surrounding sequence fixed lets us ask what the two substitutions do together.')
table(['Version','Sequence change','Measured value'],[['WT','Reference at both positions','y(WT)'],['A','Only the first variant','y(A)'],['B','Only the second variant','y(B)'],['AB','Both variants','y(AB)']],[.7,3.8,2.1])
p('The y values are on a log2 scale relative to WT. Thus y(WT) is defined as zero; this is a choice of reference, not a missing activity measurement. A value of +1 corresponds to twice the reference activity. An additive prediction on the log2 scale corresponds to multiplying the separate fold changes on the original activity scale.')
h('Define the quantity we want to predict')
p('Our experimental interaction is epsilon = y(A) + y(B) − y(WT) − y(AB). It is the expected additive activity minus the observed double-mutant activity. The code checks that this matches the negative of the source interaction coefficient, so the sign convention is explicit.')
p('As an illustrative example, suppose y(WT)=0, y(A)=0.4, and y(B)=0.3. The additive prediction for AB is 0.7. If y(AB)=0.2, epsilon is +0.5: the double has less activity than expected. If y(AB)=1.0, epsilon is −0.3: it has more activity than expected. These numbers illustrate the definition; they are not an observed case from our experiment.')
p('An interaction of zero means the two effects add on this chosen scale. Positive and negative signs describe departures from that expectation. They do not, by themselves, mean harmful, beneficial, or disease-causing.')
p('Siraj and colleagues recode alleles from lowest to highest activity and treat the lowest-activity diplotype as the reference. The pipeline applies that convention as a final step. When the least active of the four measured versions is a single mutant, A or B, the sign of epsilon is flipped; when it is WT or AB, epsilon is unchanged. The magnitude of epsilon is untouched. The identical per-pair flip is applied to the model interaction, so the two quantities always describe the same four sequences in the same order. The flip changed the sign for 1,422 of the 2,833 pairs. After recoding, mean epsilon is +0.180 and 75.9% of pairs have positive epsilon, compared with a mean of −0.002 and 51.2% positive before recoding. The noise ceiling reported below is estimated on the unrecoded contrast, which the results file keeps as epsilon_refalt.')
p('The recoding builds in a positive shift. The lowest of the four measured activities always enters the recoded epsilon with a negative sign, so measurement noise alone pushes recoded epsilon upward. In a simulation with no true interaction, with y(WT) fixed at zero and the other three activities drawn as independent noise, recoded epsilon is positive in 83% of draws. The observed 75.9% therefore cannot be read as a measured prevalence of interference, and any constant predictor gains from the shift. The convention is the published one, and it makes our interfering fraction among the 58 flagged pairs, 86.2%, comparable with the paper’s 77.2%, but the constant baselines below must be read with the shift in mind.')

page('How we selected and scored the sequences')
h('A deliberately limited dataset')
p('We selected one cell type, K562, and one released 200-base window configuration: the middle window with center_variant set to var1. Using one configuration kept the assay context consistent and avoided treating overlapping versions of the same region as independent examples. It was a scope choice, not evidence that the middle window is biologically optimal. Alternate windows and longer context remain untested sensitivities.')
p('We selected the lexicographically first library before quality filtering, retained complete quartets of single-base substitutions, required mean DNA counts of at least 20 and activity standard errors no greater than 0.5 for every haplotype, and checked the alleles and overlapping oligo sequences. We did not select pairs because their interactions were significant or because Evo scored them well. The resulting dataset has 2,833 pairs in 2,251 groups of overlapping genomic regions. The precise filtering rules and excluded rows are retained in the repository. [5]')
h('Give Evo the same four sequences')
p('We used the frozen evo2_7b_base checkpoint. Frozen means its learned weights were not updated using these MPRA measurements. For each DNA sequence, the scorer sums the log probabilities assigned to its 200 bases. Each base is predicted from earlier bases, using a beginning-of-sequence token for the initial prediction. We score the forward sequence and its reverse complement separately and average their sums.')
p('Call this number s(sequence). It measures sequence likelihood under Evo, not reporter expression. Applying the same four-way comparison gives the model interaction m = s(A) + s(B) − s(WT) − s(AB). This asks whether the model’s likelihood responds non-additively to the two changes. The signs follow the experimental convention, but matching the algebra does not guarantee that a positive model interaction corresponds to a positive biological interaction.')
h('The complete path from input to result')
p('Released measurements and oligos → validated quartets → 10,856 unique DNA sequences → cached Evo scores → one model interaction per pair → comparison with the measured interaction for that same pair.')
p('We cached each unique sequence once, then joined the scores back to the quartets by sequence hashes. The saved outputs include the per-sequence scores, per-pair predictions, summary metrics, plots, selected examples, and provenance records. This lets us revise the statistical analysis without repeating GPU inference. Today’s run used the 200-base fragments, not Evo’s maximum possible sequence context.')

page('How we evaluated the MVP')
h('First ask whether the two quantities move together')
p('The primary Spearman correlation is between the ranking of Evo interaction m and the ranking of measured experimental interaction epsilon across the 2,833 pairs. A large positive value would mean that pairs with higher model interaction tend to have higher experimental interaction. Zero means little monotonic association. Pearson correlation compares the numerical values of those same two quantities and asks whether they have a linear association.')
p('For uncertainty, we resampled overlapping-region groups 1,000 times and recalculated the statistics. This cluster bootstrap keeps related sequences together. Its 95% interval describes uncertainty under this sampling procedure; it does not cover every source of measurement bias or model error.')
h('Then ask whether the score improves numerical prediction')
p('This question is about the numerical size of the interaction, not only its rank. Evo produces a raw interaction score in log-likelihood units, while the experiment reports epsilon in log2-activity units. We therefore fit a straight line that converts the Evo score into a predicted epsilon. The line is fitted using training pairs’ Evo scores and measured epsilons, then applied to a held-out fold. This is out-of-fold calibration: each pair is predicted by a line that never used that pair or its overlapping region.')
p('A fold is simply one part of the dataset. In each of five rounds, four-fifths of the genomic groups are the training data and the remaining one-fifth is the held-out data. We learn the line from the training groups, use it to predict the held-out groups, and then move to the next round. At the end, every pair has one prediction made without using its own group’s measurement.')
p('Root mean squared error, or RMSE, summarizes the distance between each predicted epsilon and its measured epsilon. Lower RMSE means smaller prediction mistakes, and the units are log2-activity units. We compare calibrated Evo with two constant predictions: zero interaction for every pair, and the average epsilon in the training fold for every pair. These are reference strategies, not claims that every pair truly has zero or exactly the average interaction.')
h('Correcting the original baseline')
p('The original ridge regression used measured single-variant effects, including their sum, to predict epsilon. That sum also appears inside epsilon itself. Shared measured components and shared noise can generate correlation, so its apparent strength was not a clean sequence-based biological benchmark. We removed it from the primary comparison.')
p('Its replacement is a supervised ridge regression using only sequence features: counts of DNA words of lengths one, two, and three in each quartet member, plus variant positions and separation. Ridge is a linear regression with a penalty that shrinks coefficients. Scaling and fitting occur inside the training folds. This baseline can learn from training labels but never receives a test pair’s activity measurements as features. It is a simple fixed-setting comparator, not a claim to the strongest possible sequence model.')

page('What the preliminary results mean')
table(['Prediction being tested','What is predicted for each pair','RMSE or correlation','How to read it'],[
['Raw Evo association','One raw Evo interaction score is compared with measured epsilon','Spearman 0.0205\n95% interval −0.0190 to 0.0589','Do the two rankings move together? Almost not at all.'],
['Raw Evo association','The numerical Evo score is compared with measured epsilon','Pearson 0.0030\n95% interval −0.0366 to 0.0409','Is there a linear relationship? Almost not at all.'],
['Calibrated Evo','A fold-specific line uses Evo score to predict epsilon','RMSE 0.30498','How close are Evo-based numerical predictions?'],
['Zero-interaction reference','Predict epsilon = 0 for every pair','RMSE 0.35376','How well does the simplest no-interaction guess do? Under recoding this guess is biased low.'],
['Training-fold mean reference','Predict the training-fold average epsilon for every held-out pair','RMSE 0.30478','How well does a constant average do? This is the reference to beat.'],
['Sequence-only ridge','Use DNA word counts and positions to predict epsilon directly','RMSE 0.31239\nSpearman 0.0999','Better than zero, worse than the training mean. See the caution below.']],[1.55,2.15,1.55,1.55])
p('Read the table from left to right. Each row names a different way to produce a number for a pair. Every number is then compared with that pair’s measured epsilon. Correlations describe whether predictions and measurements move together; RMSE describes how far numerical predictions are from measurements. The calibrated Evo row is the main predictive test. The zero and training-mean rows show how hard it is to improve over simple constant guesses.')
p('Across all 2,833 held-out predictions, calibrated Evo has RMSE 0.30498 and the training-fold mean reference has RMSE 0.30478. The difference is 0.00020 log2-activity units, with Evo slightly worse. Calibrated Evo is worse than the training mean in four of the five held-out folds and better by less than 0.00001 in the fifth, and the paired cluster-bootstrap interval for training-mean RMSE minus Evo RMSE is −0.00042 to −0.00002. Both beat the zero-interaction reference, RMSE 0.35376, by about 0.049. That gain comes from the fitted intercept, which absorbs the positive shift the recoding introduces, not from the Evo score: the fitted slope changes sign across folds. Thus this run provides no positive RMSE gain for Evo over a constant guess; the practical difference is very small.')
p('The central finding is that the four-way Evo likelihood difference does not usefully order the measured interactions in this dataset. A Spearman value of 0.0205 is the correlation between those two rankings; it is not a percentage of correctly predicted interactions. The interval includes zero, and the calibrated model also fails to improve prediction error over the training mean. [5]')
p('The sequence-only ridge has RMSE 0.31239 and Spearman 0.0999 against measured epsilon: better than the zero reference, worse than the training mean. Calibrated Evo beats it on RMSE, with a paired interval of 0.0031 to 0.0118, but that is a comparison between two predictors that both lose to a constant. The ridge’s rank association should not be read as an epistasis signal until it is checked. The recoding decides the sign of epsilon from which measured version is least active, activity is correlated with sequence composition (GC content alone has Spearman 0.328 with reference activity), and the ridge sees sequence composition. That check has not been done. The appropriate conclusion is unchanged: the tested Evo readout has not demonstrated useful prediction of this assay’s interaction measurements.')
h('Direction and single-variant checks')
p('Among the 1,020 pairs with |epsilon| at least 0.25, the raw model interaction has the same sign as measured epsilon in 52.45% of cases. This effect-size threshold is not a significance threshold. The percentage alone does not establish reliable direction prediction; it needs uncertainty and a fair reference comparison.')
p('We also compared each single variant’s change in Evo score, s(mutant) − s(WT), with that variant’s measured activity change, y(mutant) − y(WT). Pooling A and B gives Spearman −0.0181. Comparing the absolute sizes of those same score changes and activity changes gives Spearman −0.0059. These exploratory checks suggest that the current score is also poorly aligned with single-variant activity. They do not identify why. Repeated regions and shared quartet members mean the pooled observations are not independent.')

page('Measurement noise and what remains uncertain')
h('A noisy target makes prediction harder')
p('Each measured epsilon has a standard error, epsilon_se, which describes its estimated measurement uncertainty. A smaller standard error means a more precise estimate. Across our pairs, the variance of measured epsilon is 0.12514 and the mean squared standard error is 0.08937. Under an additive, independent measurement-error model, subtracting these estimates leaves signal variance of about 0.03578. Reliability, the estimated signal share of observed variance, is therefore about 0.286.')
p('Under that model, a perfect predictor of the underlying interaction would have an observed Pearson correlation of approximately sqrt(0.286)=0.535 with the noisy measurements. This is an estimated Pearson noise ceiling for this dataset and error model. It is not a universal ceiling for all correlations or all assays. Only 256 of 2,833 pairs, or 9.04%, exceed 1.96 times their standard error in absolute value. That is a nominal per-pair test without multiple-testing correction; it does not imply that the other 91% have no biological interaction.')
h('How to interpret the noise adjustment')
p('Dividing the upper endpoint of the raw Spearman interval, 0.05895, by 0.53468 gives 0.11025. We recorded this as an approximate noise-adjustment diagnostic. However, the classical attenuation formula applies to Pearson correlation, and our calculation holds the estimated reliability fixed. It neither supplies a general correction for ranks nor propagates uncertainty in reliability. Consequently, 0.11 is a sensitivity estimate, not a rigorous upper confidence bound on the true Spearman correlation. [3, 4]')
p('The defensible current conclusion is that the observed association is near zero, and a simple noise adjustment still suggests a small association. A formal latent-effect analysis with validated errors is needed to turn that sensitivity calculation into a confidence bound.')
h('What the audit did and did not establish')
p('The sequence relationships, score joins, and four-way arithmetic passed checks. The reconstructed experimental contrast agrees within 2.22 × 10⁻¹⁶, and the model contrast within 1.78 × 10⁻¹⁵; those residuals are numerical rounding. All-zero relative WT effects are expected. Shell syntax and missing-column errors in exploratory snippets do not change the saved scoring results.')
p('Arithmetic consistency alone cannot certify inference correctness. The run records a model revision but lacks a weight-file hash, and recording the revision did not itself pin the download. The original four-haplotype design file has not been independently confirmed. A pinned rerun, inference repeatability checks, and direct construct verification remain necessary. No cross-cell comparison, learned representation probe, SAE experiment, model intervention, or new biological validation was completed today.')

page('Part 3  Build a reliable benchmark and predictor')
p('With autonomy at Arc, I would build the project in stages with explicit decisions after each stage. The aim is to establish where predictive information comes from, then test its mechanism. The following is a proposed research program; access to staff, compute, assay platforms, and data would be arranged with the relevant teams.')
h('Stage 1  Make the evidence publication ready')
p('First, reproduce inference in a supported, pinned environment and hash the actual weights. Verify an independently selected set of quartet constructs against the original design. Test repeat runs, batch-size consistency, sequence orientation, and a small set of manually checked token probabilities. Freeze a versioned data release, scoring manifest, and one-command evaluation before expanding the model search.')
p('Recover replicate-level measurements where available. Estimate interaction uncertainty from the four activities jointly, retaining covariance caused by a shared reference. Compare independent replicate estimates, then fit a measurement-error model in which each observed epsilon is a noisy estimate of an underlying pair effect. Re-estimate reliability and model association jointly, with region-level resampling. Check interval coverage using simulated data with known interactions. This directly addresses the weakest part of today’s noise claim.')
h('Stage 2  Separate cell context from sequence context')
p('Add complete quartets from other measured cells and an independent interaction dataset, after checking comparability and licensing. For cross-cell comparisons, match the same variant pair and sequence background and estimate reliability on that matched subset. Shared DNA does not imply shared regulatory output. Compare sequence-only predictions with predictions that also receive cell information, such as cell identity or a prespecified expression profile.')
p('A low cross-cell correlation alone would not prove that a sequence model cannot predict interactions. It could reflect measurement noise, assay differences, or a mixture of shared and cell-specific effects. We would quantify each contribution rather than treat that correlation as a verdict.')
h('A fixed evaluation contract')
p('Keep overlapping regions and shared constructs together. Develop models using nested grouped validation: the inner split chooses settings and the outer split estimates performance. Reserve a final independent dataset or prospective assay whose labels remain untouched during development. Compare a fixed set of likelihood readouts, checkpoint sizes, and sequence contexts, plus sequence-only and cell-aware supervised baselines. Longer genomic context is a separate experiment because it supplies DNA absent from the 200-base reporter construct.')
p('Advance a predictor only if it improves held-out association and error over simple references and competitive baselines, with uncertainty reported. For an interaction-specific claim, it must also beat an additive model using the same inputs and comparable training data. A better predictor of total activity is not automatically a better predictor of epistasis.')

page('Recover and test information inside Evo')
h('Stage 3  Ask whether the representation is more informative than the score')
p('An internal representation is a vector of numbers the model computes while processing a sequence. It can contain information that its final likelihood score does not expose. We would extract matched vectors for WT, A, B, and AB at a prespecified set of layers and positions. In an autoregressive model, positions before both mutations cannot contain information about both changes, so alignment and causal position matter.')
p('For each representation, form the same contrast: h(A) + h(B) − h(WT) − h(AB). Train a small regularized probe to predict measured epsilon from this vector. A probe is a supervised readout, not evidence that Evo was already making that prediction. Fit all preprocessing and feature selection within training folds. Compare with raw likelihood, k-mer features, single-mutant representations, and matched random or untrained features. Test whether a contrast adds value beyond the component representations themselves.')
p('The concrete output would be a reproducible comparison of readouts across layers on untouched regions and then an independent dataset. If no representation improves on the references, we should report that boundary and reconsider the dataset or question before launching a large mechanism search.')
h('Stage 4  Turn predictive features into testable hypotheses')
p('A sparse autoencoder, or SAE, decomposes a representation into features while encouraging only a small number to be active at once. These features may be easier to interpret than the original vector, but a feature is not automatically a biological entity. Arc’s Evo work provides a starting point for such analyses; compatibility with our exact checkpoint and layer must be checked. [2]')
p('For features that predict held-out epsilon, examine matched sequence changes, motif annotations, spacing, and cell dependence. An example hypothesis might be that an interaction-sensitive feature responds to two nearby transcription-factor sites. Test it on independent sequences and controls with similar GC content, mutation type, spacing, and single-variant effects. Correct for the number of features searched and confirm candidates in data not used to choose them.')
p('Then intervene inside the model: remove or replace a candidate feature’s activation at aligned positions and measure the resulting change in the model interaction or the validated probe prediction. Use matched random-feature interventions, reconstruction controls, and checks that ordinary model behavior is not broadly damaged. A selective effect supports a causal role in the model’s computation. It still does not demonstrate that the proposed transcription factors cause the interaction in a cell.')

page('Test the biology and assemble the paper')
h('Stage 5  Run a prospective experiment')
p('Design a new MPRA panel before seeing its measurements. A planning target is 200 variant pairs, or 800 individual sequence constructs before barcodes and controls, with biological replicates across at least two relevant cell settings. The final size and replicate count should follow pilot variance estimates, assay capacity, and a power calculation. Include predicted positive and negative interactions, near-additive controls, model disagreements, and an unbiased random sample. Evaluate the random sample separately so enrichment does not inflate population-level performance.')
p('For a small subset of reproducible interactions, perturb the nominated motif, alter spacing, or perturb the implicated transcription factor. Measure whether these changes alter the interaction in the predicted direction. At selected feasible loci, test the four genotype states in endogenous genomic context, with editing and genotype confirmation. A reporter experiment supports an assay-specific mechanism; endogenous validation tests whether it also holds in chromatin.')
h('The paper would answer a sequence of questions')
table(['Figure','Evidence and question'],[
['1','Validated quartet benchmark, replicate agreement, and noise estimates. What can this dataset resolve?'],
['2','Likelihood and baseline comparisons across contexts and datasets. Where does direct scoring work or fail?'],
['3','Held-out representation and probe results. Is useful information present beyond the likelihood score?'],
['4','Feature analysis and controlled model interventions. Which computations support the predictions?'],
['5','Prospective quartet measurements and targeted perturbations. Do predictions and proposed mechanisms survive experimental tests?']],[.65,5.95])
p('The strongest paper would link generalization on unseen regions to interpretable model computations and prospective biological validation. It would distinguish a feature that predicts an assay outcome, a feature that causally affects model behavior, and a biological mechanism supported by perturbation. Those are three different claims requiring three different experiments.')
p('If likelihood fails but frozen representations succeed, the paper could establish how to recover regulatory interaction information from Evo. If sequence-only prediction fails but a cell-aware model succeeds, it could identify the importance of cellular context. If every approach fails on reliable independent measurements, a rigorous benchmark paper could define a meaningful limit. None of these outcomes guarantees publication; each needs novelty, adequate power, and comparisons with existing work.')

page('Resources milestones and source record')
h('How I would use autonomy at Arc')
p('I would own the benchmark, analysis plan, experiment prioritization, integration of results, and manuscript. I would work with Evo researchers on inference and representation access, an experimental collaborator on assay design and validation, and a quantitative collaborator on measurement uncertainty and power. Arc’s combination of computational and experimental research makes that integration a plausible institutional fit; the allocations below are requests, not commitments already made. [6]')
table(['Planning period','Deliverable and decision'],[
['Weeks 1 to 3','Pinned rerun, construct audit, replicate model, and frozen evaluation splits. Proceed when measurements and inference are credible.'],
['Weeks 4 to 8','Cross-context benchmark and grouped probe comparison. Select a readout only on development data.'],
['Weeks 9 to 12','Independent computational confirmation, feature hypotheses, and assay power plan. Freeze prospective predictions.'],
['Months 4 to 6','Prospective reporter measurements and model interventions. Start targeted biological tests for reproducible cases.'],
['Months 6 to 9','Complete feasible endogenous validation, independent replication, public release, and manuscript. Revise timing around assay feasibility.']],[1.25,5.35])
p('Compute would support a measured pilot, a bounded checkpoint/context comparison, and storage of selected activations rather than every layer at every base. Experimental resources would support oligo synthesis, barcodes, sequencing, cell culture, and focused perturbations. Model expertise would help verify scoring and SAE compatibility. Additional resources should buy better controls, independent evidence, and higher measurement precision; they should not expand an unconstrained search for a positive result.')
h('Sources and reproducibility')
p('1a. Siraj and colleagues. Functional dissection of complex trait variants at single-nucleotide resolution. Nature, 2026. https://doi.org/10.1038/s41586-026-10121-6\n1b. Associated data and code. https://doi.org/10.5281/zenodo.15297965')
p('2. Arc Institute. Evo 2 implementation and checkpoint documentation. https://github.com/ArcInstitute/evo2\nRelated interpretability work: https://arcinstitute.org/news/evo2')
p('3. Saccenti and colleagues. Corruption of the Pearson correlation coefficient by measurement error and its estimation, bias, and correction under different error models. https://pmc.ncbi.nlm.nih.gov/articles/PMC6965177/\n4. Kitagawa, Nybom, and Stuhler. Measurement error and rank correlations. https://discovery.ucl.ac.uk/id/eprint/10087171/')
p('5. Project evidence: results/evo2_7b_base/metrics.json and predictions.csv; data/quartets.csv.gz, audit.csv.gz, and provenance.json in https://github.com/soleilwizman/evalingevo. This report describes the September 5 MVP and diagnostic calculations; the noise diagnostic was added through commit d110a408, the activity-based recoding through commit 270873f, and the numbers reported here come from the rerun in commit a14cb7d.\n6. Arc Institute research model. https://arcinstitute.org/about')
# Remove inherited title rules and decorative paragraph borders.
for st in D.styles:
 for node in list(st.element.iter(qn('w:pBdr'))): node.getparent().remove(node)
for node in list(D.element.iter(qn('w:pBdr'))): node.getparent().remove(node)
# Smaller references to preserve breathing room on final page.
for pp in D.paragraphs[-5:]:
 for rr in pp.runs: rr.font.size=Pt(9)
 pp.paragraph_format.space_after=Pt(5)
D.core_properties.title='Understanding regulatory interactions with Evo 2'
D.core_properties.subject='Project proposal MVP walkthrough and research plan'
D.core_properties.author=''
out=Path(__file__).resolve().parent/'final_render4'/'Evo2_regulatory_interactions_proposal_and_MVP.docx'
out.parent.mkdir(exist_ok=True)
D.save(out)
print(out)
print('Words',sum(len(pp.text.split()) for pp in D.paragraphs))
